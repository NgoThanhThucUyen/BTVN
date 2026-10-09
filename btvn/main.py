import os
import sqlite3
from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
XLSX_FILE = os.path.join(BASE_DIR, "Copy of category_853.xlsx")
DB_FILE = os.path.join(BASE_DIR, "category_853.db")
TABLE_NAME = "products"


def guess_type(values):
    """Đoán kiểu SQLite từ các giá trị của một cột."""
    for v in values:
        if v is None:
            continue
        if isinstance(v, bool):
            return "INTEGER"
        if isinstance(v, int):
            return "INTEGER"
        if isinstance(v, float):
            return "REAL"
        return "TEXT"
    return "TEXT"


def main():
    # 1. Đọc file Excel
    wb = load_workbook(XLSX_FILE, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    headers = [str(h).strip() for h in rows[0]]
    data = rows[1:]

    # 2. Xác định kiểu dữ liệu cho từng cột
    columns = []
    for i, name in enumerate(headers):
        col_type = guess_type(r[i] for r in data)
        # Nếu cột là float nhưng toàn số nguyên (vd: 99000.0) thì dùng INTEGER
        if col_type == "REAL" and all(
            r[i] is None or float(r[i]).is_integer() for r in data
        ):
            col_type = "INTEGER"
        columns.append((name, col_type))

    # 3. Tạo bảng trong SQLite
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()

    col_defs = []
    for name, col_type in columns:
        pk = " PRIMARY KEY" if name.lower() == "id" else ""
        col_defs.append(f'"{name}" {col_type}{pk}')

    cur.execute(f'DROP TABLE IF EXISTS "{TABLE_NAME}"')
    cur.execute(f'CREATE TABLE "{TABLE_NAME}" ({", ".join(col_defs)})')

    # 4. Chèn dữ liệu (ô trống sẽ thành NULL)
    placeholders = ", ".join("?" for _ in headers)
    col_names = ", ".join(f'"{h}"' for h in headers)
    sql = f'INSERT OR REPLACE INTO "{TABLE_NAME}" ({col_names}) VALUES ({placeholders})'

    clean_data = []
    for r in data:
        if all(v is None for v in r):  # bỏ dòng trống
            continue
        row = []
        for v, (_, col_type) in zip(r, columns):
            if v is not None and col_type == "INTEGER":
                v = int(v)
            row.append(v)
        clean_data.append(row)

    cur.executemany(sql, clean_data)
    conn.commit()

    # 5. Kiểm tra kết quả
    cur.execute(f'SELECT COUNT(*) FROM "{TABLE_NAME}"')
    print(f"Đã nhập {cur.fetchone()[0]} dòng vào bảng '{TABLE_NAME}'.")
    cur.execute(f'SELECT * FROM "{TABLE_NAME}" LIMIT 3')
    for row in cur.fetchall():
        print(row)

    conn.close()


if __name__ == "__main__":
    main()